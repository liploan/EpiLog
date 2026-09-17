from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_epilog_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    BG_COLOR = RGBColor(18, 24, 38)       # Slate navy
    TEXT_MAIN = RGBColor(248, 250, 252)   # Off-white
    TEXT_MUTED = RGBColor(148, 163, 184)  # Muted slate
    ACCENT_GOLD = RGBColor(234, 179, 8)   # Editorial warm gold
    CARD_BG = RGBColor(30, 41, 59)        # Card container

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="EPILOG PITCH DECK"):
        cat_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(0.4))
        p_cat = cat_box.text_frame.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_GOLD

        title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.9), Inches(11.333), Inches(0.8))
        p_title = title_box.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_MAIN

    # SLIDE 1: Title
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    t_box = s1.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.333), Inches(3.5))
    tf1 = t_box.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "EpiLog"
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_GOLD
    p2 = tf1.add_paragraph()
    p2.text = "The AI Travel Journal & Intellectual Keepsake"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_MAIN
    p2.space_before = Pt(12)
    p3 = tf1.add_paragraph()
    p3.text = "“A picture is worth a thousand words—let them write it for you.”"
    p3.font.size = Pt(18)
    p3.font.italic = True
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(24)

    # SLIDE 2: Problem
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Problem: Meaning Gets Lost in the Camera Roll")
    cards = [
        ("The Journaling Friction", "Travelers want to record memories, but writing daily logs after 15,000 steps is exhausting. By day four, journals are abandoned."),
        ("The Photo Graveyard", "Trips turn into 600+ mixed phone snaps and DSLR shots. They remain unorganized, uncontextualized, and rarely revisited."),
        ("Loss of Cultural Context", "Months later, details fade: the name of the back-alley trattoria, the history of the facade, and what made the stop unique.")
    ]
    for i, (head, desc) in enumerate(cards):
        box = s2.shapes.add_shape(1, Inches(1.0 + (i * 3.9)), Inches(2.2), Inches(3.5), Inches(4.0))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.fill.background()
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left, tf.margin_right, tf.margin_top = Inches(0.3), Inches(0.3), Inches(0.4)
        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GOLD
        p_body = tf.add_paragraph()
        p_body.text = desc
        p_body.font.size = Pt(14)
        p_body.font.color.rgb = TEXT_MAIN
        p_body.space_before = Pt(14)

    # SLIDE 3: Solution
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "The Solution: Automated, Editorial Chronicles")
    sol_box = s3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.5))
    tf3 = sol_box.text_frame
    tf3.word_wrap = True
    sol_points = [
        ("Zero-Effort Capture", "Drop in a vacation photo folder. EpiLog parses timestamps and coordinates to reconstruct routes and timeline stops automatically."),
        ("Multimodal Scene Intelligence", "Computer vision identifies landmarks, dishes, and activities, producing concise narrative journals without manual prompts."),
        ("Reflections & 'What I Learned'", "Synthesizes architectural techniques, culinary origins, and cultural history into lasting intellectual takeaways."),
        ("Historical Time Capsule", "Integrates major world events and headlines from those exact calendar dates, providing temporal context.")
    ]
    for title, desc in sol_points:
        p = tf3.add_paragraph()
        p.text = f"• {title}: "
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GOLD
        p.space_before = Pt(14)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MAIN

    # SLIDE 4: Tech Innovation
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Core Innovation: Intelligent Sensor Fusion")
    tech_box = s4.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.5))
    tf4 = tech_box.text_frame
    tf4.word_wrap = True
    tech_steps = [
        ("1. Temporal Alignment", "Creates a unified timeline by indexing universal EXIF timestamps across phone and camera files."),
        ("2. Proximity GPS Tagging", "Matches orphan DSLR shots to phone GPS anchors taken within temporal proximity (+/- 3 mins) to inherit coordinates."),
        ("3. VLM Visual Bridging", "When time gaps exist, a vision model compares scene traits to bridge GPS-less DSLR shots to accurate locations."),
        ("4. Editorial Curation", "Automatically surfaces high-resolution DSLR photos as hero imagery while retaining phone geodata.")
    ]
    for title, desc in tech_steps:
        p = tf4.add_paragraph()
        p.text = f"{title}: "
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GOLD
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MAIN

    # SLIDE 5: Deliverables
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Product Deliverables: Screen to Print")
    deliv_cards = [
        ("Interactive Dashboard", "Split-screen web workspace featuring route maps synchronized with an editable narrative journal feed."),
        ("Social Event Cards", "1-click generation of editorial 9:16 vertical cards with route stamps, photos, and reflections for social channels."),
        ("Fine-Art Coffee Table Book", "High-margin physical export: layflat binding, archival art paper, and automated typesetting for a bespoke volume.")
    ]
    for i, (title, desc) in enumerate(deliv_cards):
        box = s5.shapes.add_shape(1, Inches(1.0 + (i * 3.9)), Inches(2.2), Inches(3.5), Inches(4.0))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.fill.background()
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left, tf.margin_right, tf.margin_top = Inches(0.3), Inches(0.3), Inches(0.4)
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GOLD
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = TEXT_MAIN
        p_desc.space_before = Pt(14)

    # SLIDE 6: Business Model
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Business Model & Monetization")
    b_box = s6.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.5))
    tf6 = b_box.text_frame
    tf6.word_wrap = True
    model_points = [
        ("Freemium Digital App", "Free core route mapping and photo clusters. Premium unlocks full VLM reflections and high-res vector exports."),
        ("High-Margin Physical Books", "Print-on-demand luxury monographs retailing at $120–$250+. Zero inventory risk with healthy unit margins."),
        ("Travel Ecosystem Partnerships", "White-label or affiliate integrations with boutique agencies, safari operators, and airlines for post-trip recaps.")
    ]
    for title, desc in model_points:
        p = tf6.add_paragraph()
        p.text = f"• {title}: "
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GOLD
        p.space_before = Pt(16)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MAIN

    # SLIDE 7: Roadmap
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Product Roadmap")
    phases = [
        ("Phase 1: Web MVP", "Browser-side EXIF parsing, MapLibre route plotting, DBSCAN stop clustering, and Gemini 1.5 Flash pipeline."),
        ("Phase 2: Desktop & Print", "Tauri desktop app for managing RAW/DSLR libraries, paired with an automated Paged.js print-ready PDF generator."),
        ("Phase 3: Mobile Native", "iOS & Android apps enabling passive background geofencing, real-time stop suggestions, and direct camera roll sync.")
    ]
    for i, (phase, desc) in enumerate(phases):
        box = s7.shapes.add_shape(1, Inches(1.0), Inches(2.2 + (i * 1.5)), Inches(11.333), Inches(1.2))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.fill.background()
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left, tf.margin_top = Inches(0.4), Inches(0.2)
        p = tf.paragraphs[0]
        p.text = phase
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GOLD
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = TEXT_MAIN
        p_desc.space_before = Pt(4)

    # SLIDE 8: Vision Close
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    c_box = s8.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.0))
    tfc = c_box.text_frame
    tfc.word_wrap = True
    pc1 = tfc.paragraphs[0]
    pc1.text = "EpiLog"
    pc1.font.size = Pt(48)
    pc1.font.bold = True
    pc1.font.color.rgb = ACCENT_GOLD
    pc2 = tfc.add_paragraph()
    pc2.text = "Turning raw travel data into lasting personal literature."
    pc2.font.size = Pt(22)
    pc2.font.bold = True
    pc2.font.color.rgb = TEXT_MAIN
    pc2.space_before = Pt(12)
    pc3 = tfc.add_paragraph()
    pc3.text = "Your journey, remembered and understood."
    pc3.font.size = Pt(16)
    pc3.font.italic = True
    pc3.font.color.rgb = TEXT_MUTED
    pc3.space_before = Pt(16)

    prs.save("EpiLog_Pitch_Deck.pptx")
    print("Successfully generated EpiLog_Pitch_Deck.pptx")

if __name__ == "__main__":
    create_epilog_deck()
